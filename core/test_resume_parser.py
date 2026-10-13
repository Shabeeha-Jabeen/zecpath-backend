
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from docx import Document

from core.resume_parser import extract_resume_text, clean_resume_text


class ResumeParserTests(SimpleTestCase):

    def test_clean_resume_text(self):
        sample_text = "  Shabeeha Jabeen  \n\n  Python Developer  "

        cleaned_text = clean_resume_text(sample_text)

        self.assertIn("Shabeeha Jabeen", cleaned_text)
        self.assertIn("Python Developer", cleaned_text)
        self.assertNotIn("  ", cleaned_text)

    def test_extract_resume_text_from_docx(self):
        document = Document()
        document.add_paragraph("Shabeeha Jabeen")
        document.add_paragraph("Python Django Developer")

        docx_buffer = BytesIO()
        document.save(docx_buffer)

        docx_file = SimpleUploadedFile(
            "sample_resume.docx",
            docx_buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

        extracted_text = extract_resume_text(docx_file)

        self.assertIn("Shabeeha Jabeen", extracted_text)
        self.assertIn("Python Django Developer", extracted_text)
